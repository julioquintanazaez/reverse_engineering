
import pandas as pd
from io import BytesIO
from django.http import HttpResponse
from collections import defaultdict
import openpyxl
from openpyxl import load_workbook, Workbook

def json_to_excel(json_data):
    # Convertir el defaultdict a dict normal si es necesario
    if isinstance(json_data, defaultdict):
        json_data = dict(json_data)
    
    pacientes = list(json_data[next(iter(json_data))].keys())
    
    # Separar campos en diferentes categorías
    info_personal = ['Accession number', 'First name', 'Last name', 'Middle initial', 
                    'Ordering physician', 'DOB', 'SSN', 'Gender', 'Ethnicity',
                    'Collection date', 'Received date', 'Report generated',
                    'Specimen type', 'Report format type', 'Current medications']
    
    info_medica = ['MEDICATIONS THAT HAVE BEEN PROBLEMATIC',
                  'MEDICATION ALLERGIES', 'BRIEF MEDICAL HISTORY', 'Renal function',
                  'Smoker', 'Daily or nearly daily alcohol intake', 'Language']
    
    # Crear archivo Excel en memoria usando openpyxl directamente
    output = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Genotype Results'
    
    # Escribir información personal sin encabezados
    row_num = 1
    for campo in info_personal:
        if campo in json_data:
            worksheet.cell(row=row_num, column=1, value=campo)
            for col_num, paciente in enumerate(pacientes, start=2):
                worksheet.cell(row=row_num, column=col_num, value=json_data[campo].get(paciente, ''))
            row_num += 1
    
    # Añadir espacio entre secciones
    row_num += 2
    
    # Escribir información médica sin encabezados
    for campo in info_medica:
        if campo in json_data:
            worksheet.cell(row=row_num, column=1, value=campo)
            for col_num, paciente in enumerate(pacientes, start=2):
                worksheet.cell(row=row_num, column=col_num, value=json_data[campo].get(paciente, ''))
            row_num += 1
    
    # Añadir espacio antes de SNPs
    row_num += 2
    
    # Escribir encabezado para SNPs (personalizado)
    worksheet.cell(row=row_num, column=1, value='rs#')
    for col_num, paciente in enumerate(pacientes, start=2):
        worksheet.cell(row=row_num-1, column=col_num, value=f'Patient {col_num-1}')
        worksheet.cell(row=row_num, column=col_num, value='Genotype')
    row_num += 1
    
    # Escribir datos de SNPs
    rs_fields = [k for k in json_data.keys() if k not in info_personal + info_medica and k != 'Patient ID']
    for campo in rs_fields:
        worksheet.cell(row=row_num, column=1, value=campo)
        for col_num, paciente in enumerate(pacientes, start=2):
            worksheet.cell(row=row_num, column=col_num, value=json_data[campo].get(paciente, ''))
        row_num += 1
    
    # Ajustar anchos de columnas
    for col in worksheet.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        worksheet.column_dimensions[column_letter].width = adjusted_width
    
    workbook.save(output)
    output.seek(0)
    return output